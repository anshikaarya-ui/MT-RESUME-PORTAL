from __future__ import annotations

from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

from backend.utils import clean_text


class ExtractionError(Exception):
    pass


def extract_text_from_file(path: str | Path) -> str:
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(file_path)
    if suffix == ".docx":
        return _extract_docx(file_path)
    if suffix == ".txt":
        return _extract_txt(file_path)

    raise ExtractionError(f"Unsupported file type: {suffix}")


def _extract_pdf(path: Path) -> str:
    try:
        chunks: list[str] = []
        with fitz.open(path) as pdf:
            for page in pdf:
                chunks.append(page.get_text("text"))
        text = clean_text("\n".join(chunks))
        if len(text) < 80:
            raise ExtractionError(
                "Could not extract enough text from this PDF. It may be scanned; add OCR for production."
            )
        return text
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"PDF extraction failed: {exc}") from exc


def _extract_docx(path: Path) -> str:
    try:
        doc = Document(path)
        paragraphs = [p.text for p in doc.paragraphs]
        table_text: list[str] = []
        for table in doc.tables:
            for row in table.rows:
                table_text.append(" | ".join(cell.text for cell in row.cells))
        text = clean_text("\n".join(paragraphs + table_text))
        if len(text) < 50:
            raise ExtractionError("Could not extract enough text from this DOCX.")
        return text
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"DOCX extraction failed: {exc}") from exc


def _extract_txt(path: Path) -> str:
    try:
        return clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    except Exception as exc:
        raise ExtractionError(f"TXT extraction failed: {exc}") from exc
